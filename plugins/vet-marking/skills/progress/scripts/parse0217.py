"""Parse a Wisenet Report 0217 'Unit Enrolment Outcome Matrix' (.xls/.xlsx) into tidy rows.
Usage: python3 -I parse0217.py <report.xls|xlsx> <outdir>
Writes <outdir>/sheets.csv (inventory), learners.csv, outcomes.csv, units.csv
outcomes.csv has black=True where the cell is blacked out. RTO rule: black = not resulted yet (outstanding), never excluded.
"""
import sys, os, re, subprocess, datetime as dt, csv
import openpyxl

src, out = sys.argv[1], sys.argv[2]
os.makedirs(out, exist_ok=True)
if src.lower().endswith('.xls'):
    subprocess.run(['soffice','--headless','--convert-to','xlsx','--outdir',out,src],
                   check=True, capture_output=True, timeout=180)
    src = os.path.join(out, os.path.splitext(os.path.basename(src))[0] + '.xlsx')
wb = openpyxl.load_workbook(src)  # not read_only: need merged ranges + hidden state

UNIT_RE = re.compile(r'^\s*([A-Z]{3,}[A-Z0-9]*\d{3,4}[A-Z]?)\s*\n?\s*(\d{2}/\d{2}/\d{4})\s*-\s*(\d{2}/\d{2}/\d{4})', re.S)
def d(s): return dt.datetime.strptime(s, '%d/%m/%Y').date()
def asdate(v):
    if isinstance(v, dt.datetime): return v.date()
    if isinstance(v, dt.date): return v
    if isinstance(v, str):
        for f in ('%d/%m/%Y','%Y-%m-%d'):
            try: return dt.datetime.strptime(v.strip()[:10], f).date()
            except ValueError: pass
    return None

def is_black(cell):
    """True for a black-filled cell. RTO rule: black = not resulted yet."""
    f = cell.fill
    if not f or f.fill_type != 'solid': return False
    rgb = str(f.fgColor.rgb or '')
    if len(rgb) == 8 and rgb[2:].isalnum():
        try:
            r, g, b = (int(rgb[i:i+2], 16) for i in (2, 4, 6))
            return 0.2126*r + 0.7152*g + 0.0722*b < 50
        except ValueError: return False
    return False

def find(ws, label):
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.strip().lower().startswith(label.lower()):
                return c
    return None

def right_value(ws, cell):
    for col in range(cell.column + 1, ws.max_column + 1):
        v = ws.cell(cell.row, col).value
        if v not in (None, ''): return v
    return None

inv, learners, outcomes, units = [], [], [], []
for ws in wb.worksheets:
    rec = dict(sheet=ws.title, state=ws.sheet_state, rows=ws.max_row, cols=ws.max_column)
    title = find(ws, 'Unit Enrolment Outcome Matrix')
    ucell, lcell = find(ws, 'Units'), find(ws, 'Learner Name')
    if not (title and ucell and lcell):
        rec['type'] = 'not an outcome matrix - review manually'; inv.append(rec); continue
    meta = {}
    for lab, key in [('Course Offer Code', 'course_code'), ('Course Offer Desc', 'course_desc'),
                     ('Location', 'location'), ('Offer Start Date', 'offer_start')]:
        c = find(ws, lab); meta[key] = right_value(ws, c) if c else None
    meta['offer_start'] = asdate(meta['offer_start'])
    # unit columns: header cells right of 'Units'; a merged header covers every column in its range
    merged = {}
    for m in ws.merged_cells.ranges:
        if m.min_row <= ucell.row <= m.max_row:
            for col in range(m.min_col, m.max_col + 1): merged[col] = (m.min_col, m.max_col)
    ucols = []
    for col in range(ucell.column + 1, ws.max_column + 1):
        v = ws.cell(ucell.row, col).value
        if isinstance(v, str) and UNIT_RE.match(v):
            code, s, e = UNIT_RE.match(v).groups()
            lo, hi = merged.get(col, (col, col))
            ucols.append((code, d(s), d(e), list(range(lo, hi + 1))))
            units.append(dict(sheet=ws.title, unit=code, unit_start=d(s), unit_end=d(e), col=col))
    # learner columns by header label
    hdr = {}
    for col in range(1, ws.max_column + 1):
        v = ws.cell(lcell.row, col).value
        if isinstance(v, str) and v.strip(): hdr[v.strip()] = col
    rs = hdr.get('Reg Start / End Date'); ref = hdr.get('RefInternal (RefExternal)')
    st = hdr.get('Status'); wp = hdr.get('Workplace')
    n = 0
    for r in range(lcell.row + 1, ws.max_row + 1):
        name = ws.cell(r, lcell.column).value
        if name in (None, '') or not str(name).strip(): continue
        start = asdate(ws.cell(r, rs).value) if rs else None
        end = None
        if rs:
            for col in range(rs + 1, (ref or rs + 4)):
                end = asdate(ws.cell(r, col).value) or end
        refv = str(ws.cell(r, ref).value or '').strip() if ref else ''
        m = re.match(r'^(\S*)\s*\((.*)\)', refv)
        sid, ext = (m.group(1), m.group(2).strip()) if m else (refv, '')
        L = dict(sheet=ws.title, **meta, learner=str(name).strip(), student_id=sid, ref_external=ext,
                 reg_start=start, reg_end=end, status=(ws.cell(r, st).value if st else None),
                 workplace=(ws.cell(r, wp).value if wp else None), row=r)
        learners.append(L); n += 1
        for code, us, ue, cols in ucols:
            black = any(is_black(ws.cell(r, c)) for c in cols)
            vals = [ws.cell(r, c).value for c in cols if ws.cell(r, c).value not in (None, '')]
            oc = str(vals[0]).strip() if vals else ''
            if oc.endswith('.0'): oc = oc[:-2]
            outcomes.append(dict(sheet=ws.title, student_id=sid, learner=L['learner'], unit=code,
                                 unit_start=us, unit_end=ue, outcome=oc, black=black, conflict=len(set(map(str, vals))) > 1))
    rec.update(type='outcome matrix', header_row=lcell.row, learners=n, units=len(ucols), **meta)
    inv.append(rec)

def dump(name, rows):
    if not rows: return
    keys = list(dict.fromkeys(k for r in rows for k in r))
    with open(os.path.join(out, name), 'w', newline='') as f:
        w = csv.DictWriter(f, keys); w.writeheader(); w.writerows(rows)
for nm, rows in [('sheets.csv', inv), ('learners.csv', learners), ('outcomes.csv', outcomes), ('units.csv', units)]:
    dump(nm, rows)
for r in inv: print(r)
